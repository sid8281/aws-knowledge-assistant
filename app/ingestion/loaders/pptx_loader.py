"""Extract text from a PowerPoint (.pptx) file using python-pptx."""

from pptx import Presentation


def load_pptx(file_path: str) -> str:
    prs = Presentation(file_path)
    slide_texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                slide_texts.append(shape.text_frame.text)
    return "\n".join(slide_texts)
